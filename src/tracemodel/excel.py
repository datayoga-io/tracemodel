"""A model as an Excel workbook, for people who read models in Excel: model.to_excel(path).

Sheets: Summary (counts, check status, legend, links); the grids, the model's items laid out
as a finance team reads them; then the model as tables: Dimensions, Items, Checks, Check
results and Values (every entered value with its source).

In a grid, items are blocks one under the other, sharing the columns (the period, or another
dim); each block has its own rows, its other dims nested in the first column (a hierarchy
innermost, its tree under each label of the others), indented and grouped.
`sections` names groups of items, e.g. a statement each; with layout="sheet" they all go on
one sheet, one after the other, each under a banner and on its own printed page.

Cells follow the modelling convention: blue is entered, black is calculated; a yellow fill
is a value entered over a formula, a red fill a cell a failing check reads (or that it found
missing). Nil shows as "—"; a cell with no value is blank. Each cell's note gives its source
or its formula.

Needs openpyxl: pip install pytracemodel[excel].
"""

from __future__ import annotations

import contextlib
import itertools
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .errors import DefinitionError
from .expr import Cell
from .values import NIL

if TYPE_CHECKING:
    from .model import Item, Model
    from .values import Report

LAYOUTS = ("sheet", "sheets")
_TABLE_SHEETS = ("Dimensions", "Items", "Checks", "Check results", "Values")


def to_excel(
    model: Model,
    path: str | Path,
    *,
    title: str = "Model",
    sections: Mapping[str, Sequence[str]] | None = None,
    layout: str = "sheet",
    columns: str | None = None,
    split_after: tuple[str, str] | None = None,
    page_breaks: bool = True,
) -> None:
    """Write the model to an .xlsx workbook. See the module docstring for what it holds."""
    xl = _openpyxl()
    if layout not in LAYOUTS:
        raise DefinitionError(f"to_excel: layout must be one of {', '.join(map(repr, LAYOUTS))}, got {layout!r}")
    if columns is not None and columns not in model.dimensions:
        raise DefinitionError(f"to_excel: columns {columns!r} is not a dimension")
    groups = _groups(model, sections or {})
    sheets = _sheets(model, groups, layout, columns, sections is not None)

    report = model.check()
    w = _Writer(xl, model, report, columns, split_after, page_breaks)
    wb = xl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    taken = {"Summary", *_TABLE_SHEETS}
    sheet_of: dict[str, str] = {}
    grids = []
    for name, (across, secs) in sheets.items():
        ws = wb.create_sheet(_sheet_name(name, taken))
        items = [i for _, its in secs for i in its]
        grids.append((ws, items))
        sheet_of |= {i.name: ws.title for i in items}
        w.grid(ws, name, across, secs)
    w.dimensions(wb.create_sheet("Dimensions"))
    w.items(wb.create_sheet("Items"), sheet_of)
    w.checks(wb.create_sheet("Checks"))
    w.results(wb.create_sheet("Check results"))
    w.values(wb.create_sheet("Values"))
    w.summary(summary, title, grids)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        if ws.title not in _TABLE_SHEETS:  # grids print one page wide; tables keep their size
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(path)


def _openpyxl():
    try:
        import openpyxl
        from openpyxl import comments, styles, utils
        from openpyxl.worksheet import pagebreak, table
    except ImportError:
        raise ImportError("to_excel needs openpyxl: pip install pytracemodel[excel]") from None
    ns = type("openpyxl", (), {})
    ns.Workbook = openpyxl.Workbook
    ns.Comment = comments.Comment
    ns.Alignment, ns.Border, ns.Font = styles.Alignment, styles.Border, styles.Font
    ns.PatternFill, ns.Side = styles.PatternFill, styles.Side
    ns.get_column_letter = utils.get_column_letter
    ns.Break = pagebreak.Break
    ns.Table, ns.TableStyleInfo = table.Table, table.TableStyleInfo
    return ns


# == which items go where ====================================================

def _groups(model: Model, sections: Mapping[str, Sequence[str]]) -> list[tuple[str, list[Item]]]:
    """The named sections, then the other items grouped by their dims."""
    out, placed = [], set()
    for name, item_names in sections.items():
        items = [model.item(n) for n in item_names]
        if not items:
            raise DefinitionError(f"to_excel: section {name!r} has no items")
        for n in item_names:
            if n in placed:
                raise DefinitionError(f"to_excel: item {n!r} is in two sections")
            placed.add(n)
        out.append((name, items))
    by_dims: dict[tuple[str, ...], list[Item]] = {}
    for item in model.items.values():
        if item.name not in placed:
            by_dims.setdefault(tuple(sorted(item.dims)), []).append(item)
    for items in by_dims.values():
        dims = items[0].dims
        name = _title(items[0]) if len(items) == 1 else f"by {' × '.join(dims)}" if dims else "inputs"
        out.append((name, items))
    return out


def _sheets(model, groups, layout, columns, named) -> dict[str, tuple[str | None, list]]:
    """{sheet name: (the dim across, its sections)}. With layout="sheet", the sections that
    share the dim across go on one sheet; the others keep a sheet each."""
    sheets: dict[str, tuple[str | None, list]] = {}
    shared = []
    for name, items in groups:
        across = _across(model, items[0].dims, columns)
        if across is not None and not all(across in i.dims for i in items):
            across = None
        if layout == "sheet" and across is not None and (not shared or shared[0][0] == across):
            shared.append((across, (name, items)))
        else:
            sheets[name] = (across, [(name, items)])
    if len(shared) == 1:
        across, (name, items) = shared[0]
        sheets = {name: (across, [(name, items)])} | sheets
    elif shared:
        sheets = {"Statements" if named else "Model": (shared[0][0], [s for _, s in shared])} | sheets
    return sheets


def _across(model, dims, columns) -> str | None:
    """The dim across: the one asked for, else period, else the last dim that isn't a hierarchy
    (hierarchies read down); None when every dim is a hierarchy."""
    if columns in dims:
        return columns
    if "period" in dims:
        return "period"
    flat = [d for d in dims if not model.dimensions[d].hierarchy]
    return flat[-1] if flat else None


# == the workbook ============================================================

class _Writer:
    def __init__(self, xl, model: Model, report: Report, columns, split_after, page_breaks):
        self.xl, self.model, self.report = xl, model, report
        self.columns, self.split_after, self.page_breaks = columns, split_after, page_breaks
        self.flagged = _flagged(model, report)
        F, P, S = xl.Font, xl.PatternFill, xl.Side
        self.BLUE, self.BLUE_BOLD = F(color="0000FF"), F(color="0000FF", bold=True)        # entered
        self.BLACK, self.BLACK_BOLD = F(color="000000"), F(color="000000", bold=True)      # calculated
        self.GREY, self.BOLD, self.TITLE = F(color="808080"), F(bold=True), F(bold=True, size=14)
        self.LINK = F(color="0563C1", underline="single")
        self.OVERRIDE = P("solid", fgColor="FFF2CC")   # entered over a formula
        self.FAILING = P("solid", fgColor="F8CBAD")    # read by a failing check
        self.HEADER = P("solid", fgColor="DDEBF7")
        self.BANNER, self.BANNER_FONT = P("solid", fgColor="1F4E78"), F(bold=True, size=12, color="FFFFFF")
        self.EDGE, self.THIN = S(style="medium", color="404040"), S(style="thin")
        amber = P("solid", fgColor="FFEB9C")
        self.STATUS = {"pass": P("solid", fgColor="C6EFCE"), "FAIL": P("solid", fgColor="FFC7CE"),
                       "missing": amber, "warn": amber, "warn: missing": amber}

    # -- grids ---------------------------------------------------------------
    def grid(self, ws, name, across, sections) -> None:
        items = [i for _, its in sections for i in its]
        if across is None and len(items) > 1 and all(i.dims == items[0].dims for i in items):
            return self.items_across(ws, name, items)
        xl, model = self.xl, self.model
        col_labels = model.dimensions[across].coords if across else ("value",)
        stacked, sectioned = len(items) > 1, len(sections) > 1
        downs = {i.name: _nesting(model, [d for d in i.dims if d != across]) for i in items}
        n_labels = 1  # the dims down share the first column, nested (a pivot table's compact form)
        units_col = n_labels + 1
        first = units_col + 1
        ws.sheet_view.showGridLines = False

        ws.cell(1, 1, name if stacked else _title(items[0])).font = self.TITLE
        ws.cell(2, 1, f"{len(sections)} sections, {len(items)} items by {across}, one after the other" if sectioned
                else f"{len(items)} items by {across or 'their dims'}, one block each" if stacked
                else " × ".join(items[0].dims)).font = self.GREY
        r = header_row = 4
        same = [downs[i.name] for i in items]
        same_dims = all(d == same[0] for d in same)
        self._head(ws.cell(r, 1, " / ".join(same[0]) if same_dims else ""))
        self._head(ws.cell(r, units_col, "units"))
        for j, label in enumerate(col_labels):
            c = ws.cell(r, first + j, label)
            self._head(c)
            c.alignment = xl.Alignment(horizontal="right")
        split_col = None
        if self.split_after and across == self.split_after[0] and self.split_after[1] in col_labels:
            split_col = first + col_labels.index(self.split_after[1])
        ws.freeze_panes = ws.cell(r + 1, first)
        ws.print_title_rows = f"{r}:{r}"  # the column header on every printed page
        r += 1

        def values(row, item, fmt, at, bold):
            for j, label in enumerate(col_labels):
                self._value(ws.cell(row, first + j), item, at | ({across: label} if across else {}), fmt, bold)
                if split_col == first + j:
                    ws.cell(row, first + j).border = xl.Border(right=self.EDGE)

        base = 1 if sectioned else 0  # outline level of a block's title row
        for title, its in sections:
            if sectioned:  # the section's banner, across the whole width; a new printed page
                if self.page_breaks and r > header_row + 1:
                    ws.row_breaks.append(xl.Break(id=r - 1))
                for col in range(1, first + len(col_labels)):
                    ws.cell(r, col).fill = self.BANNER
                ws.cell(r, 1, title).font = self.BANNER_FONT
                ws.row_dimensions[r].height = 20
                r += 1
            for item in its:
                down = downs[item.name]
                fmt = _number_format(model, item)
                if stacked:  # the block's title row; an item with no other dims has its values here
                    c = ws.cell(r, 1, _title(item))
                    c.font = self.BOLD
                    c.alignment = xl.Alignment(horizontal="left", indent=base)
                    c.comment = xl.Comment(f"item: {item.name}", "tracemodel", width=240, height=40)
                    ws.cell(r, units_col, item.unit_label).font = self.GREY
                    if base:
                        ws.row_dimensions[r].outline_level = base
                    if not down:
                        values(r, item, fmt, {}, False)
                        r += 1
                        continue
                    r += 1
                r = self._rows(ws, r, item, down, across, col_labels, values, fmt, base + stacked, units_col)
                if stacked:
                    if base:
                        ws.row_dimensions[r].outline_level = base
                    r += 1  # a blank row between blocks
        ws.sheet_properties.outlinePr.summaryBelow = False  # parents sit above their children
        widths = {1: 58}
        widths[units_col] = 18
        widths |= {first + j: 14 for j in range(len(col_labels))}
        self._widths(ws, widths)

    def _rows(self, ws, r, item, down, across, col_labels, values, fmt, indent, units_col) -> int:
        """An item's rows, its dims down nested in the first column: each label of an outer dim is
        a group's row, with the next dim's rows beneath it; the last dim's rows hold the values.
        A group with no values is left out (a grid of several dims is sparse)."""
        model, xl = self.model, self.xl
        if not down:  # an item with only the dim across: one row of values
            if not indent:
                ws.cell(r, units_col, item.unit_label).font = self.GREY
            values(r, item, fmt, {}, False)
            return r + 1
        *outer, inner = down

        def label_row(r, d, label, level):
            dim = model.dimensions[d]
            depth = _depth(dim, label) + level + indent
            c = ws.cell(r, 1, dim.description(label) or label)
            c.alignment = xl.Alignment(horizontal="left", indent=depth)
            parent = level < len(outer) or bool(dim.hierarchy and dim.hierarchy.children(label))
            if parent:
                c.font = self.BOLD
            if depth:
                ws.row_dimensions[r].outline_level = min(depth, 7)
            return parent

        def has_values(at, rest):
            cols = col_labels if across else (None,)
            return any(_state(model, item, at | dict(zip(rest, key)) | ({across: lb} if across else {}))[0]
                       is not None for key in itertools.product(*(model.dimensions[d].coords for d in rest))
                       for lb in cols)

        def nest(r, at, level):
            if level == len(outer):
                for label in model.dimensions[inner].coords:
                    bold = label_row(r, inner, label, level)
                    if not indent:
                        ws.cell(r, units_col, item.unit_label).font = self.GREY
                    values(r, item, fmt, at | {inner: label}, bold)
                    r += 1
                return r
            d = outer[level]
            for label in model.dimensions[d].coords:
                sub = at | {d: label}
                if has_values(sub, down[level + 1:]):
                    label_row(r, d, label, level)
                    r = nest(r + 1, sub, level + 1)
            return r

        return nest(r, {}, 0)

    def items_across(self, ws, name, items) -> None:
        """Items that share only hierarchy dims: the items across (the item name as a dim)."""
        xl, model, dims = self.xl, self.model, items[0].dims
        ws.sheet_view.showGridLines = False
        ws.cell(1, 1, name).font = self.TITLE
        ws.cell(2, 1, f"{' × '.join(dims)} down, items across").font = self.GREY
        r = 4
        for i, d in enumerate(dims, start=1):
            self._head(ws.cell(r, i, d))
        for j, item in enumerate(items):
            c = ws.cell(r, len(dims) + 1 + j, f"{_title(item)}\n{item.unit_label}".strip())
            self._head(c)
            c.alignment = xl.Alignment(horizontal="right", wrap_text=True)
        ws.row_dimensions[r].height = 30
        ws.freeze_panes = ws.cell(r + 1, len(dims) + 1)
        ws.print_title_rows = f"{r}:{r}"
        fmts = [_number_format(model, i) for i in items]
        for key in itertools.product(*(model.dimensions[d].coords for d in dims)):
            r += 1
            at = dict(zip(dims, key))
            parent = False
            for i, (d, label) in enumerate(zip(dims, key), start=1):
                dim = model.dimensions[d]
                c = ws.cell(r, i, dim.description(label) or label)
                c.alignment = xl.Alignment(horizontal="left", indent=_depth(dim, label))
                if dim.hierarchy and dim.hierarchy.children(label):
                    c.font, parent = self.BOLD, True
            if len(dims) == 1 and _depth(model.dimensions[dims[0]], key[0]):
                ws.row_dimensions[r].outline_level = min(_depth(model.dimensions[dims[0]], key[0]), 7)
            for j, item in enumerate(items):
                self._value(ws.cell(r, len(dims) + 1 + j), item, at, fmts[j], parent)
        ws.sheet_properties.outlinePr.summaryBelow = False
        self._widths(ws, {**{i: 58 if len(dims) == 1 else 30 for i in range(1, len(dims) + 1)},
                          **{len(dims) + 1 + j: 18 for j in range(len(items))}})

    def _value(self, c, item, full, fmt, bold) -> None:
        xl = self.xl
        value, kind, note, key = _state(self.model, item, full)
        if value is NIL:
            c.value, c.font = "—", self.GREY
            c.alignment = xl.Alignment(horizontal="right")
        elif value is not None:
            c.value = value
            entered = kind in ("entered", "override", "cell_formula")
            c.font = (self.BLUE_BOLD if bold else self.BLUE) if entered else (self.BLACK_BOLD if bold else self.BLACK)
            if kind == "override":
                c.fill = self.OVERRIDE
            c.number_format = ("yyyy-mm-dd hh:mm" if isinstance(value, datetime)
                               else "yyyy-mm-dd" if isinstance(value, date) else fmt)
            _as_text(c)
        if note:
            c.comment = xl.Comment(note, "tracemodel", width=320, height=90)
        if key in self.flagged:
            c.fill = self.FAILING

    # -- the model as tables ---------------------------------------------------
    def summary(self, ws, title, grids) -> None:
        model, report, xl = self.model, self.report, self.xl
        ws["A1"] = title
        ws["A1"].font = self.TITLE
        n_pass = sum(r.passed for r in report.results)
        status = f"{n_pass}/{len(report.results)} runs pass"
        if report.errors:
            status += f", {len(report.errors)} errors"
        if report.warnings:
            status += f", {len(report.warnings)} warnings"
        rows = [("Dimensions", len(model.dimensions)), ("Items", len(model.items)), ("Values entered", len(model)),
                ("Checks", f"{len(model.checks)} checks, {len(report.results)} runs"), ("Check status", status)]
        for i, (k, v) in enumerate(rows, start=3):
            ws.cell(i, 1, k).font = self.BOLD
            ws.cell(i, 2, v)
        r = 4 + len(rows)
        ws.cell(r, 1, "Legend").font = self.BOLD
        for sample, font, fill, text in [
            ("1,234", self.BLUE, None, "Entered, with a source (in the cell's note, and on the Values sheet)"),
            ("1,234", self.BLACK, None, "Calculated by a formula (in the cell's note)"),
            ("1,234", self.BLUE, self.OVERRIDE, "Entered over a formula: overrides it for this cell"),
            ("—", self.GREY, None, "Nil: the source shows nothing here; counts as 0"),
            ("", None, None, "Blank: not entered"),
            ("1,234", self.BLUE, self.FAILING, "Read by a failing check, or missing for one (see Check results)"),
        ]:
            r += 1
            c = ws.cell(r, 1, sample)
            c.alignment = xl.Alignment(horizontal="right")
            if font:
                c.font = font
            if fill:
                c.fill = fill
            ws.cell(r, 2, text)
        r += 2
        for col, head in enumerate(("Sheet", "Items", "Dims"), start=1):
            ws.cell(r, col, head).font = self.BOLD
        for gs, items in grids:
            r += 1
            self._link(ws.cell(r, 1, gs.title), gs.title)
            ws.cell(r, 2, ", ".join(i.name for i in items))
            ws.cell(r, 3, " × ".join(items[0].dims) or "(none)")
        for sheet in _TABLE_SHEETS:
            r += 1
            self._link(ws.cell(r, 1, sheet), sheet)
        self._widths(ws, {1: 22, 2: 70, 3: 40})

    def dimensions(self, ws) -> None:
        dims = self.model.dimensions.values()
        keys: list[str] = []
        for d in dims:
            for label in d.coords:
                keys += [k for k in d.attrs(label) if k not in keys]
        rows = [[d.name, label, d.description(label), d.hierarchy.parent(label) if d.hierarchy else None,
                 _depth(d, label), *_dates(d, label), *(d.attrs(label).get(k) for k in keys)]
                for d in dims for label in d.coords]
        header = ["dimension", "label", "description", "parent", "depth", "start", "end", *keys]
        self._table(ws, "Dimensions", header, rows, {1: 22, 2: 44, 3: 50, 6: 12, 7: 12})
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
            for c in row:
                c.number_format = "yyyy-mm-dd"

    def items(self, ws, sheet_of) -> None:
        items = list(self.model.items.values())
        keys: list[str] = []
        for i in items:
            keys += [k for k in i.attrs if k not in keys and k not in ("units", "scale")]
        header = ["item", "description", "dims", "dtype", "units", "scale", "formula", "sheet", *keys]
        rows = [[i.name, i.description, " × ".join(i.dims), i.dtype.__name__, i.attrs.get("units"),
                 i.attrs.get("scale"), repr(i.formula) if i.formula is not None else None, sheet_of.get(i.name),
                 *(i.attrs.get(k) for k in keys)] for i in items]
        self._table(ws, "Items", header, rows, {1: 32, 2: 44, 3: 30, 7: 40, 8: 22})
        for r, i in enumerate(items, start=2):
            if i.name in sheet_of:
                self._link(ws.cell(r, 8), sheet_of[i.name])

    def checks(self, ws) -> None:
        rows = []
        for c in self.model.checks.values():
            expect = repr(c.expect)
            rows.append([c.id, c.severity, c.description, ", ".join(c.coords) or "-", expect])
        self._table(ws, "Checks", ["check", "severity", "description", "runs over", "expect"], rows,
                    {1: 22, 3: 60, 4: 24, 5: 90})

    def results(self, ws) -> None:
        from .values import _status

        rows = [[r.check.id, ", ".join(f"{d}={v}" for d, v in r.at.items()), _status(r), r.lhs,
                 getattr(r.expect, "rule", ""), r.rhs, r.diff, "; ".join(str(c) for c in r.cells), r.error]
                for r in self.report.results]
        header = ["check", "coords", "result", "left", "rule", "right", "diff", "cells read", "error"]
        self._table(ws, "CheckResults", header, rows, {1: 22, 2: 30, 8: 100, 9: 40})
        for i, row in enumerate(rows, start=2):
            ws.cell(i, 3).fill = self.STATUS.get(row[2], self.STATUS["FAIL"])

    def values(self, ws) -> None:
        from .model import _is_formula

        dims = list(self.model.dimensions)
        header = ["item", *dims, "value", "kind", "document", "section", "page", "quote", "url", "overrides_formula"]
        rows = []
        for e in self.model:
            kind = "nil" if e.value is NIL else "formula" if _is_formula(e.value) else "value"
            shown = f"= {e.value!r}" if kind == "formula" else e.value
            s = e.source
            rows.append([e.item, *(e.coords.get(d) for d in dims), shown, kind, s.get("document"),
                         s.get("section"), s.get("page"), s.get("quote"), s.get("url"), e.overrides_formula or None])
        n = len(dims)
        self._table(ws, "Values", header, rows, {1: 32, n + 2: 16, n + 4: 40, n + 5: 50})
        for i, row in enumerate(rows, start=2):
            if row[n + 7]:
                c = ws.cell(i, n + 8)
                c.hyperlink, c.font = row[n + 7], self.LINK

    # -- helpers -------------------------------------------------------------
    def _table(self, ws, name, header, rows, widths) -> None:
        xl = self.xl
        ws.append(header)
        for row in rows:
            ws.append([_plain(v) for v in row])
            for c in ws[ws.max_row]:
                _as_text(c)
        if rows:  # an Excel table needs a data row
            t = xl.Table(displayName=name, ref=f"A1:{xl.get_column_letter(len(header))}{len(rows) + 1}")
            t.tableStyleInfo = xl.TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(t)
        else:
            for c in ws[1]:
                self._head(c)
        ws.freeze_panes = "B2"
        self._widths(ws, {i: widths.get(i, 14) for i in range(1, len(header) + 1)})

    def _head(self, c) -> None:
        c.font, c.fill, c.border = self.BOLD, self.HEADER, self.xl.Border(bottom=self.THIN)

    def _link(self, c, sheet) -> None:
        c.hyperlink, c.font = f"#'{sheet}'!A1", self.LINK

    def _widths(self, ws, widths) -> None:
        for i, w in widths.items():
            ws.column_dimensions[self.xl.get_column_letter(i)].width = w


# == cells ===================================================================

def _state(model: Model, item: Item, full: Mapping[str, str]) -> tuple[Any, str | None, str | None, tuple]:
    """(value, kind, note, key): kind is entered, override, cell_formula or calculated; None
    with no value."""
    from .model import _is_formula

    key = (item.name, tuple(full[d] for d in item.dims))
    e = model._entries.get(key)
    if e is None and item.formula is None:
        return None, None, None, key
    value = model._value(Cell(model, item, dict(full)), {})
    if e is not None and _is_formula(e.value):
        return value, "cell_formula", f"= {e.value!r}\nSource: {_source(e.source)}", key
    if e is not None:
        note = f"Source: {_source(e.source)}"
        if e.overrides_formula:
            return value, "override", f"{note}\nOverrides the formula: {item.formula!r}", key
        return value, "entered", note, key
    return value, "calculated", f"= {item.formula!r}", key


def _flagged(model: Model, report: Report) -> set:
    """Cells to mark: what a failing run reads directly (not its formulas' inputs), and any
    cell it found missing."""
    from .cellfn import run_check
    from .formula import Formula

    out = set()
    for r in report.failures:
        out |= {(c.item, tuple(c.coords.values())) for c in r.missing}
        reads: list = []
        if isinstance(r.check.expect, Formula):
            with contextlib.suppress(Exception):  # only the reads matter
                run_check(r.check.expect, model, r.at, reads)
        for c, at in reads:
            with contextlib.suppress(Exception):  # a read that doesn't resolve marks nothing
                full = c.resolve(at)
                out.add((c.item.name, tuple(full[d] for d in c.item.dims)))
    return out


def _number_format(model: Model, item: Item) -> str:
    """Whole numbers, negatives in parentheses; two decimals when the item has fractions."""
    whole = "#,##0;(#,##0)"
    if item.dtype is str:
        return "@"
    if item.dtype is not float:
        return whole
    for key in itertools.product(*(model.dimensions[d].coords for d in item.dims)):
        v = _state(model, item, dict(zip(item.dims, key)))[0]
        if isinstance(v, float) and abs(v - round(v)) > 1e-9:
            return "#,##0.00###;(#,##0.00###)"
    return whole


def _dates(dim, label) -> tuple[date | None, date | None]:
    """A time dimension's label: its first and last day."""
    if dim.time is None:
        return None, None
    p = dim.period(label)
    return p.start_time, p.end_time


def _title(item: Item) -> str:
    """The item as people read it: its description, else its name in words."""
    return item.description or item.name.replace("_", " ").capitalize()


def _source(s: Mapping[str, Any]) -> str:
    text = ", ".join(str(s[k]) for k in ("document", "section") if s.get(k))
    return text + (f", p. {s['page']}" if s.get("page") else "")


def _nesting(model: Model, down: list[str]) -> list[str]:
    """The dims down, outermost first: a hierarchy reads as the rows themselves, so it goes
    innermost and its tree stays whole under each label of the others."""
    tree = next((d for d in down if model.dimensions[d].hierarchy), None)
    return [d for d in down if d != tree] + [tree] if tree else down


def _depth(dim, label) -> int:
    if not dim.hierarchy:
        return 0
    depth, p = 0, dim.hierarchy.parent(label)
    while p is not None:
        depth, p = depth + 1, dim.hierarchy.parent(p)
    return depth


def _plain(v: Any) -> Any:
    """A value a cell can hold."""
    if v is None or isinstance(v, (int, float, str, date)):
        return v
    return "—" if v is NIL else str(v)


def _as_text(c) -> None:
    """Text stays text: openpyxl writes a string starting with "=" as an Excel formula."""
    if isinstance(c.value, str) and c.data_type == "f":
        c.data_type = "s"


def _sheet_name(name: str, taken: set[str]) -> str:
    """A unique sheet name Excel accepts: 31 characters, none of []:*?/\\."""
    base = "".join("-" if ch in "[]:*?/\\" else ch for ch in name)[:31].strip("'") or "Sheet"
    out, n = base, 2
    while out in taken:
        out = f"{base[:28]}~{n}"
        n += 1
    taken.add(out)
    return out
